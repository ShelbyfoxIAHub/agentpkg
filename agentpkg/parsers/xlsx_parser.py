"""XLSX text extraction."""

from pathlib import Path
from io import BytesIO


def parse_xlsx(path: Path) -> str:
    """Return non-empty cell values from every worksheet in an XLSX file."""
    return parse_xlsx_bytes(Path(path).read_bytes())


def parse_xlsx_bytes(data: bytes) -> str:
    """Return non-empty cell values from XLSX bytes."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "XLSX parsing requires the optional dependencies; install "
            "'agentpkg[documents]'."
        ) from error

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    rows.append("\t".join(values))
            if rows:
                sheets.append(
                    "{}\n{}".format(worksheet.title, "\n".join(rows))
                )
        return "\n\n".join(sheets)
    finally:
        workbook.close()
