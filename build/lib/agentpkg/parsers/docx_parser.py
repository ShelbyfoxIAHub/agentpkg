"""DOCX text extraction."""

from pathlib import Path
from io import BytesIO


def parse_docx(path: Path) -> str:
    """Return paragraph and table text from a DOCX file."""
    return parse_docx_bytes(Path(path).read_bytes())


def parse_docx_bytes(data: bytes) -> str:
    """Return paragraph and table text from DOCX bytes."""
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError(
            "DOCX parsing requires the optional dependencies; install "
            "'agentpkg[documents]'."
        ) from error

    document = Document(BytesIO(data))
    sections = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            sections.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(sections)
