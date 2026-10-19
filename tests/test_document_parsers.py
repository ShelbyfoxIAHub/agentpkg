import importlib.util
import unittest
from io import BytesIO

from agentpkg.parsers.docx_parser import parse_docx_bytes
from agentpkg.parsers.pdf_parser import parse_pdf_bytes
from agentpkg.parsers.xlsx_parser import parse_xlsx_bytes


class OptionalDocumentParserTests(unittest.TestCase):
    @unittest.skipUnless(
        importlib.util.find_spec("pypdf"),
        "Install agentpkg[documents] to test PDF parsing.",
    )
    def test_pdf_parser_accepts_pdf_bytes(self):
        content = b"BT /F1 12 Tf 72 720 Td (Agent package PDF) Tj ET"
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            (
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
            ),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            (
                b"<< /Length "
                + str(len(content)).encode("ascii")
                + b" >>\nstream\n"
                + content
                + b"\nendstream"
            ),
        ]
        pdf = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for index, obj in enumerate(objects, start=1):
            offsets.append(len(pdf))
            pdf.extend(
                str(index).encode("ascii")
                + b" 0 obj\n"
                + obj
                + b"\nendobj\n"
            )
        xref_offset = len(pdf)
        pdf.extend(b"xref\n0 6\n0000000000 65535 f \n")
        for offset in offsets[1:]:
            pdf.extend("{:010d} 00000 n \n".format(offset).encode("ascii"))
        pdf.extend(
            b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n"
            + str(xref_offset).encode("ascii")
            + b"\n%%EOF"
        )

        self.assertEqual(parse_pdf_bytes(bytes(pdf)), "Agent package PDF")

    @unittest.skipUnless(
        importlib.util.find_spec("docx"),
        "Install agentpkg[documents] to test DOCX parsing.",
    )
    def test_docx_parser_extracts_paragraph_and_table_text(self):
        from docx import Document

        output = BytesIO()
        document = Document()
        document.add_paragraph("Agent instructions")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "tool"
        table.cell(0, 1).text = "search"
        document.save(output)

        self.assertEqual(
            parse_docx_bytes(output.getvalue()),
            "Agent instructions\ntool\tsearch",
        )

    @unittest.skipUnless(
        importlib.util.find_spec("openpyxl"),
        "Install agentpkg[documents] to test XLSX parsing.",
    )
    def test_xlsx_parser_extracts_worksheet_values(self):
        from openpyxl import Workbook

        output = BytesIO()
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Vectors"
        worksheet.append(["dimension", "value"])
        worksheet.append([1, 0.25])
        workbook.save(output)

        self.assertEqual(
            parse_xlsx_bytes(output.getvalue()),
            "Vectors\ndimension\tvalue\n1\t0.25",
        )


if __name__ == "__main__":
    unittest.main()
